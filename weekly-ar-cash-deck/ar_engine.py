#!/usr/bin/env python3
"""
ar_engine.py — aggregates a raw AR line list (sheet 'AR' of the weekly AR extract) with the AR dashboard rules.

Rules (validated against the AR dashboard and the BU meeting Power BI report):
  * Advance  = Status 1 'Advance' lines, excluding J-6-* (negative advance lines) and C-3-a
  * Unbilled = Status 1 'Unbilled' lines, excluding 'J-6-a Unearned on LS'
  * Unpaid   = Status 1 'Unpaid' lines
  * Amounts  = 'Total Amount HT', in MMAD
  * Ownership: At risk = C-1-a/C-1-b, SS postponed = J-3-c, Under JESA = other J-* codes, Under OCP = C-* codes and PDR
  * Collections (payments received) = Status 1 'Paid' lines by 'Payment Date', on the 'Amount Invoiced' basis
    (the basis used for 'Invoices paid in <month>' in the BU meeting report); labor / non-labor from 'PSR/PROC'.
"""
import collections, datetime as dt, warnings
import openpyxl

warnings.filterwarnings("ignore")

BU_CANON = {
    "advisory": "Advisory", "ads": "Advisory",
    "am": "AM", "amt": "AM",
    "building & capabilities": "Building & Capabilities", "b&c": "Building & Capabilities",
    "building & infrastructure": "Building & Infrastructure", "bin": "Building & Infrastructure", "b&i": "Building & Infrastructure",
    "energy operations": "Energy Operations", "ene": "Energy Operations", "energy": "Energy Operations",
    "fertilizers & chemicals": "Fertilizers & Chemicals", "fec": "Fertilizers & Chemicals",
    "global mining": "Global Mining", "gmm": "Global Mining",
    "jesa technologies": "JESA Technologies", "jet": "JESA Technologies",
    "sp2m ad hoc program": "SP2M Ad Hoc Program", "sp2m": "SP2M Ad Hoc Program",
    "water & environnement": "Water & Environment", "water & environement": "Water & Environment",
    "water & environment": "Water & Environment", "wen": "Water & Environment",
    "power to x": "Power to X",
}
BU_ORDER = ["Water & Environment", "Fertilizers & Chemicals", "SP2M Ad Hoc Program", "AM", "Energy Operations",
            "Global Mining", "Building & Infrastructure", "JESA Technologies", "Advisory", "Building & Capabilities", "Power to X"]

def canon_bu(name):
    if name is None: return None
    k = " ".join(str(name).strip().lower().split())
    if k in BU_CANON: return BU_CANON[k]
    for key, v in BU_CANON.items():
        if k.startswith(key): return v
    return None

def s2_code(label):
    """'C-5-a Signed by …' / 'C-5.a Signed by …' -> 'c-5.a'; 'PDR' -> 'pdr'."""
    l = str(label or "").strip()
    if len(l) >= 5 and l[1] == "-" and l[3] in "-.": return (l[:3] + "." + l[4]).lower()
    return l.lower()[:12]

STATUS_LABELS = {  # canonical labels for the coded statuses
    "c-1.a": "At risk (new project)", "c-1.b": "At risk (budget overrun)", "c-1.c": "Pending client PO", "c-1.d": "Pending PO correction",
    "c-1.e": "Waiting mirror POs for PDR", "c-2.a": "Waiting OCP VAT exemption", "c-3.a": "Pending DP payment / clearing",
    "c-3.b": "Invoicing waiting DP clearing", "c-4.a": "Pending client PM approval", "c-5.a": "Signed by client PM, pending OCP SS approval",
    "c-5.b": "Not signed by client PM, pending OCP SS approval", "c-6.a": "Pending SS OCP rejection", "c-7.a": "Current", "c-7.b": "Past due",
    "j-1.a": "Pending JESA DMR / attachment creation", "j-1.b": "Pending attachments JESA PA", "j-1.c": "Pending LS milestone",
    "j-1.d": "Under JESA finance review", "j-1.e": "Under JESA PC review", "j-1.f": "Pending LC by PC", "j-2.a": "Pending JESA PM approval",
    "j-3.a": "Signed by client, pending JESA SS submission", "j-3.c": "Service sheets postponed", "j-4.a": "Ongoing invoicing",
    "j-5.a": "Rejected by client", "j-6.a": "Unearned on LS", "j-7.a": "Under investigation", "pdr": "PDR",
}
def status_label(code):
    c = code[:3].upper() + code[3:] if len(code) >= 5 and code[1] == "-" else code.upper()
    return f"{c} {STATUS_LABELS[code]}" if code in STATUS_LABELS else c

def owner_of(code):
    if code in ("c-1.a", "c-1.b"): return "At risk"
    if code == "j-3.c": return "SS postponed"
    if code.startswith("j-"): return "Under JESA"
    return "Under OCP"

def unpaid_phase(comment):
    c = str(comment or "").strip().lower()
    if "hors ocp" in c: return "Non-OCP / OCP sub (hors)"
    if "ocp sub" in c: return "OCP subs"
    if c.startswith("rejet"): return "Rejected"
    if c.startswith("transmise"): return "Treasury"
    if c.startswith("en ordonnancement"): return "Ordonnancement"
    if c.startswith("manque"): return "Missing reception"
    if c.startswith("en cours"): return "Being processed"
    return "Other"

def unpaid_group(code):
    return {"c-7.a": "Current", "c-7.b": "Past due", "j-5.a": "Rejected"}.get(code, "Other")

def bucket(label):
    l = str(label or "").strip().lower()
    if l.startswith("less than 30"): return "< 30 days"
    if l in ("31 to 60", "30 to 60"): return "31–60 days"
    if l in ("61 to 90", "60 to 90"): return "61–90 days"
    if l in ("91 to 120", "90 to 120"): return "91–120 days"
    if l.startswith("over 120"): return "> 120 days"
    return "n/a"
BUCKETS = ["< 30 days", "31–60 days", "61–90 days", "91–120 days", "> 120 days", "n/a"]

def num(v):
    """MAD -> MMAD"""
    return float(v) / 1e6 if isinstance(v, (int, float)) else 0.0

def aggregate(path, as_of, prev_cutoff=None):
    """Aggregate a raw AR line list. as_of: date of the snapshot (for the current month). prev_cutoff: date (payments up to
    that date count as 'previous week' collections). Returns a dict of plain-python structures (MMAD)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["AR"] if "AR" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h else "" for h in next(it)]
    I = {h: i for i, h in enumerate(hdr)}
    def g(r, name):
        i = I.get(name); return r[i] if i is not None else None
    D = lambda: collections.defaultdict(float)
    kpi = D(); bu = D(); owner = D(); st_tot = D(); st_bu = D(); upg = D(); uph = D(); agu = D(); agp = D(); year = D(); labor = D(); ocp = D()
    proj = D(); adv_bu = D(); st_proj = {}
    nonocp = D(); nonocp_proj = D(); nonocp_meta = {}; jesa_unb = D(); jesa_unb_st = D()
    coll_month = D(); coll_bu = D(); coll_cnt = collections.Counter(); dep_bu = D(); dep_tot = 0.0
    cur_m = (as_of.year, as_of.month)
    prior_m = ((as_of.replace(day=1) - dt.timedelta(days=1)).year, (as_of.replace(day=1) - dt.timedelta(days=1)).month)
    for r in it:
        if not any(v is not None for v in r): continue
        s1 = str(g(r, "Status 1") or "").strip().lower()
        amt = num(g(r, "Total Amount HT")); code = s2_code(g(r, "Status 2"))
        b = canon_bu(g(r, "Business Unit")) or (str(g(r, "Business Unit")).strip() if g(r, "Business Unit") else "Other")
        # ---- collections (paid lines)
        if s1.startswith("paid"):
            pd = g(r, "Payment Date")
            if isinstance(pd, (dt.datetime, dt.date)):
                pdd = pd.date() if isinstance(pd, dt.datetime) else pd
                inv = num(g(r, "Amount Invoiced")); lab = "labor" if str(g(r, "PSR/PROC") or "").strip().upper() == "PSR" else "non_labor"
                if pdd.year == as_of.year:
                    coll_month[(pdd.month, lab)] += inv; coll_bu[(b, "ytd")] += inv; coll_cnt[b] += 1
                    if prev_cutoff and pdd <= prev_cutoff: coll_bu[(b, "ytd_prev")] += inv
                    if (pdd.year, pdd.month) == cur_m:
                        coll_bu[(b, "month")] += inv
                        if prev_cutoff and pdd <= prev_cutoff: coll_bu[(b, "month_prev")] += inv
                    if (pdd.year, pdd.month) == prior_m: coll_bu[(b, "prior")] += inv
            rd = g(r, "Reception date")
            if isinstance(rd, (dt.datetime, dt.date)) and (rd.year, rd.month) == cur_m: dep_bu[b] += amt; dep_tot += amt
            continue
        comp = "Advance" if s1.startswith("advance") else ("Unbilled" if s1.startswith("unbilled") else ("Unpaid" if s1.startswith("unpaid") else None))
        rd = g(r, "Reception date")
        if comp == "Unpaid" and isinstance(rd, (dt.datetime, dt.date)) and (rd.year, rd.month) == cur_m: dep_bu[b] += amt; dep_tot += amt
        if comp is None: continue
        if comp == "Unbilled" and code == "j-6.a": continue  # unearned on LS is not unbilled AR; C-3.a (DP clearing) stays in unbilled, as in the BU meeting report
        if comp == "Advance" and (code.startswith("j-6") or code == "c-3.a"): continue
        kpi[comp] += amt; bu[(b, comp)] += amt
        y = str(g(r, "Year") or "").strip(); year[y if y[:2] == "20" else "No year (advance lines)"] += amt
        labor[(comp, "labor" if str(g(r, "PSR/PROC") or "").strip().upper() == "PSR" else "non_labor")] += amt
        c1 = str(g(r, "Client 1") or "").strip().upper(); is_non = c1.startswith("NON OCP"); ocp["NON OCP" if is_non else "OCP"] += amt
        pname = str(g(r, "Project Name") or "").strip(); wbs = str(g(r, "New WBS") or g(r, "Project") or "").strip()
        lab = "labor" if str(g(r, "PSR/PROC") or "").strip().upper() == "PSR" else "non_labor"
        if is_non and comp in ("Unbilled", "Unpaid"):
            nonocp[(b, comp)] += amt; key = (b, wbs, pname)
            nonocp_proj[(key, comp)] += amt; nonocp_meta.setdefault(key, (str(g(r, "Project Client") or "").strip(), status_label(code)))
        if comp == "Unbilled" and owner_of(code) == "Under JESA":
            jesa_unb[(b, lab)] += amt; jesa_unb_st[(code, lab)] += amt
        if comp == "Unbilled":
            owner[(b, owner_of(code))] += amt; st_tot[code] += amt; st_bu[(b, code)] += amt
            agu[(b, bucket(g(r, "Category Unbilled")))] += amt
            proj[(b, "Unbilled", wbs, pname)] += amt; st_proj.setdefault((b, "Unbilled", wbs, pname), code)
        elif comp == "Unpaid":
            upg[(b, unpaid_group(code))] += amt; uph[(b, unpaid_phase(g(r, "OCP Comments")))] += amt
            agp[(b, bucket(g(r, "Category")))] += amt
            proj[(b, "Unpaid", wbs, pname)] += amt; st_proj.setdefault((b, "Unpaid", wbs, pname), code)
        else:
            adv_bu[(b, status_label(code))] += amt
    kpi["Total AR"] = kpi["Advance"] + kpi["Unbilled"] + kpi["Unpaid"]
    # top projects per BU
    top = collections.defaultdict(list)
    for (b, t, wbs, pname), v in proj.items():
        top[(b, t)].append((wbs, pname, v, status_label(st_proj[(b, t, wbs, pname)])))
    top = {k: sorted(v, key=lambda x: -x[2])[:8] for k, v in top.items()}
    nproj = []
    for key, (client, st) in nonocp_meta.items():
        nproj.append([key[0], client, key[1], key[2], nonocp_proj.get((key, "Unbilled"), 0.0), nonocp_proj.get((key, "Unpaid"), 0.0), st])
    nproj.sort(key=lambda r: -(r[4] + r[5]))
    return {"kpi": dict(kpi), "bu": dict(bu), "owner": dict(owner), "status_total": dict(st_tot), "status_bu": dict(st_bu),
            "unpaid_group": dict(upg), "unpaid_phase": dict(uph), "aging_unbilled": dict(agu), "aging_unpaid": dict(agp), "year": dict(year),
            "labor": dict(labor), "ocp": dict(ocp), "advance_status": dict(adv_bu), "top_projects": top,
            "nonocp": dict(nonocp), "nonocp_projects": nproj, "jesa_unbilled": dict(jesa_unb), "jesa_unbilled_status": dict(jesa_unb_st),
            "coll_month": dict(coll_month), "coll_bu": dict(coll_bu), "coll_cnt": dict(coll_cnt), "deposited_bu": dict(dep_bu), "deposited_total": dep_tot}

def bus_in(agg):
    return [b for b in BU_ORDER if any(k[0] == b for k in agg["bu"])] + sorted({k[0] for k in agg["bu"]} - set(BU_ORDER))
